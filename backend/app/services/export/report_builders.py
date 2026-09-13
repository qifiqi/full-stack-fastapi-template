"""Word/Excel export builders (P6: python-docx + openpyxl pipelines)."""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from openpyxl import Workbook


def build_word_report(
    *,
    task: Any,
    summary: dict[str, Any],
    best_rows: list[dict[str, Any]],
) -> bytes:
    """Strategy backtest report: task header + summary table + best grid."""
    document = Document()
    document.add_heading(f"策略回测报告 - {task.name}", level=0)
    document.add_paragraph(f"任务 ID: {task.id}")
    document.add_paragraph(f"任务类型: {task.task_type}")
    document.add_paragraph(f"状态: {task.status}")

    document.add_heading("汇总", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    header = table.rows[0].cells
    header[0].text = "指标"
    header[1].text = "数值"
    for key in ("total", "success_count", "failed_count"):
        row = table.add_row().cells
        row[0].text = key
        row[1].text = str(summary.get(key, ""))

    document.add_heading("最优结果", level=1)
    grid = document.add_table(rows=1, cols=4)
    grid.style = "Light Grid Accent 1"
    cells = grid.rows[0].cells
    cells[0].text = "股票"
    cells[1].text = "模型"
    cells[2].text = "周期"
    cells[3].text = "最优指标"
    for item in best_rows:
        cells = grid.add_row().cells
        cells[0].text = str(item.get("stock_code") or "")
        cells[1].text = str(item.get("model_key") or "")
        cells[2].text = str(item.get("period_key") or "")
        cells[
            3
        ].text = f"{item.get('best_metric_name')}: {item.get('best_metric_value')}"

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_excel_summary(rows: list[dict[str, Any]]) -> bytes:
    """Model-summary grid as an xlsx workbook."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "model_summary"
    if rows:
        columns = list(rows[0].keys())
        sheet.append(columns)
        for row in rows:
            sheet.append([row.get(c) for c in columns])
    else:
        sheet.append(["empty"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
